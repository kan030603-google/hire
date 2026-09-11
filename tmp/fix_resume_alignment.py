from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


BODY_WIDTHS_MM = (128.0, 32.0)
COMPANY_ROW_SIZE_PT = 10.0


def get_or_add(parent, tag: str):
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    return element


def lock_two_column_table(table) -> None:
    widths_twips = [int(Mm(width).twips) for width in BODY_WIDTHS_MM]
    grid_columns = list(table._tbl.tblGrid.gridCol_lst)
    if len(grid_columns) != 2:
        raise ValueError("Expected a two-column table")

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    tbl_pr = table._tbl.tblPr
    tbl_w = get_or_add(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = get_or_add(tbl_pr, "w:tblInd")
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = get_or_add(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    for grid_column, width_twips in zip(grid_columns, widths_twips):
        grid_column.set(qn("w:w"), str(width_twips))

    for row in table.rows:
        cells = row.cells
        if len(cells) != 2:
            raise ValueError("Unexpected row structure in two-column table")
        for cell, width_twips in zip(cells, widths_twips):
            tc_w = get_or_add(cell._tc.get_or_add_tcPr(), "w:tcW")
            tc_w.set(qn("w:w"), str(width_twips))
            tc_w.set(qn("w:type"), "dxa")


def fix_alignment(input_path: Path, output_path: Path) -> None:
    document = Document(input_path)
    fixed = 0
    company_rows = 0
    for table in document.tables:
        if len(table._tbl.tblGrid.gridCol_lst) == 2:
            lock_two_column_table(table)
            fixed += 1
            for row in table.rows:
                if "实习生" in row.cells[0].text:
                    for run in row.cells[0].paragraphs[0].runs:
                        run.font.size = Pt(COMPANY_ROW_SIZE_PT)
                    company_rows += 1

    if fixed != 5:
        raise ValueError(f"Expected 5 body tables, found {fixed}")
    if company_rows != 2:
        raise ValueError(f"Expected 2 company rows, found {company_rows}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: fix_resume_alignment.py <input.docx> <output.docx>")
    fix_alignment(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
