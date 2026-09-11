from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


FONT_EAST_ASIA = "Microsoft YaHei"
FONT_LATIN = "Arial"
INK = "1F2328"
MUTED = "666666"
RULE = "555555"
PLACEHOLDER = "A6A6A6"


def set_run_font(run, size: float, *, bold: bool = False, color: str = INK) -> None:
    run.font.name = FONT_LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    r_fonts.set(qn("w:cs"), FONT_LATIN)


def set_style_font(style, size: float, *, bold: bool = False, color: str = INK) -> None:
    style.font.name = FONT_LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    r_fonts.set(qn("w:cs"), FONT_LATIN)


def set_cell_width(cell, width_mm: float) -> None:
    cell.width = Mm(width_mm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Mm(width_mm).twips)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, *, top=0, start=0, bottom=0, end=0) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = f"w:{side}"
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(int(value)))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, *, color: str = PLACEHOLDER, size: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "start", "bottom", "end"):
        tag = f"w:{edge}"
        node = tc_borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_bottom_border(cell, *, color: str = RULE, size: int = 10) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    bottom = tc_borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        tc_borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), color)


def set_raw_tc_border(tc, edge: str, *, color: str, size: int) -> None:
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    tag = f"w:{edge}"
    node = tc_borders.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        tc_borders.append(node)
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:space"), "0")
    node.set(qn("w:color"), color)


def set_table_fixed(table) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def reset_cell_to_one_paragraph(cell):
    for paragraph in list(cell.paragraphs[1:]):
        paragraph._element.getparent().remove(paragraph._element)
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    return paragraph


def format_paragraph(
    paragraph,
    *,
    before: float = 0,
    after: float = 0,
    line: float | None = None,
    keep_with_next: bool | None = None,
    keep_together: bool | None = None,
) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = Pt(line)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if keep_with_next is not None:
        pf.keep_with_next = keep_with_next
    if keep_together is not None:
        pf.keep_together = keep_together
    pf.widow_control = True


def add_bottom_border(paragraph, *, color: str = RULE, size: int = 10, space: int = 3) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)


def add_section_heading(doc: Document, text: str, *, compact: bool = False) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    run = paragraph.add_run(text)
    set_run_font(run, 14.4, bold=True)
    format_paragraph(
        paragraph,
        before=4.0 if compact else 6.5,
        after=4.5,
        line=17.8,
        keep_with_next=True,
        keep_together=True,
    )
    add_bottom_border(paragraph)


def add_labelled_text(paragraph, text: str, *, size: float = 9.25, bold_label: bool = True) -> None:
    match = re.match(r"^([^：:]+[：:])(.*)$", text)
    if not match:
        set_run_font(paragraph.add_run(text), size)
        return
    label, rest = match.groups()
    set_run_font(paragraph.add_run(label), size, bold=bold_label)
    set_run_font(paragraph.add_run(rest), size)


def add_bullet(doc: Document, text: str, *, size: float = 10.3, after: float = 1.8) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    pf = paragraph.paragraph_format
    pf.left_indent = Mm(7.0)
    pf.first_line_indent = Mm(-3.4)
    pf.right_indent = Mm(4.0)
    format_paragraph(paragraph, after=after, line=15.6, keep_together=True)
    add_labelled_text(paragraph, text, size=size)


def add_header_block(doc: Document, name: str, contact: str, first_heading: str) -> None:
    table = doc.add_table(rows=2, cols=3)
    set_table_fixed(table)
    widths = (26.0, 108.0, 26.0)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell, top=0, start=25, bottom=0, end=25)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    top_row = table.rows[0]
    heading_row = table.rows[1]
    top_row.height = Mm(22.0)
    top_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    heading_row.height = Mm(8.0)
    heading_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    clear_paragraph(top_row.cells[0].paragraphs[0])

    center = top_row.cells[1]
    name_p = center.paragraphs[0]
    name_p.style = doc.styles["Title"]
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    display_name = re.sub(r"\s+", "", name)
    set_run_font(name_p.add_run(display_name), 24.5, bold=True)
    format_paragraph(name_p, after=1.5, line=28.5, keep_with_next=True)

    contact_p = center.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parts = [part.strip() for part in re.split(r"[｜|]", contact) if part.strip()]
    if len(parts) >= 2:
        set_run_font(contact_p.add_run("手机号码："), 9.6, bold=True)
        set_run_font(contact_p.add_run(parts[0]), 9.6)
        set_run_font(contact_p.add_run("  |  "), 9.6, color=MUTED)
        set_run_font(contact_p.add_run("邮箱："), 9.6, bold=True)
        set_run_font(contact_p.add_run(parts[1]), 9.6)
    else:
        set_run_font(contact_p.add_run(contact), 9.6)
    format_paragraph(contact_p, line=13.0)

    heading_cell = heading_row.cells[0].merge(heading_row.cells[1])
    photo = top_row.cells[2].merge(heading_row.cells[2])
    set_cell_borders(photo)
    set_cell_bottom_border(photo)
    # Word renders the continuation cell's edges for a vertical merge. Apply
    # them directly so the photo placeholder remains a complete rectangle.
    photo_continuation_tc = heading_row._tr.tc_lst[-1]
    set_raw_tc_border(photo_continuation_tc, "start", color=PLACEHOLDER, size=8)
    set_raw_tc_border(photo_continuation_tc, "end", color=PLACEHOLDER, size=8)
    set_raw_tc_border(photo_continuation_tc, "bottom", color=RULE, size=10)
    photo_p = reset_cell_to_one_paragraph(photo)
    photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(photo_p.add_run("照片位置"), 8.0, color=MUTED)
    format_paragraph(photo_p, line=11.0)

    set_cell_margins(heading_cell, top=0, start=0, bottom=0, end=0)
    set_cell_bottom_border(heading_cell)
    heading_p = reset_cell_to_one_paragraph(heading_cell)
    heading_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(heading_p.add_run(first_heading), 14.2, bold=True)
    format_paragraph(heading_p, line=17.2, keep_with_next=True, keep_together=True)


def split_row(text: str) -> tuple[str, str]:
    if "\t" in text:
        left, right = text.rsplit("\t", 1)
        return left.strip(), right.strip()
    match = re.match(r"^(.*?)(\d{4}\.\d{2}\s*-\s*\d{4}\.\d{2})$", text)
    if match:
        return match.group(1).strip(), match.group(2).strip()
    return text.strip(), ""


def add_education_row(doc: Document, text: str) -> None:
    left, date = split_row(text)
    parts = [part.strip() for part in left.split("｜") if part.strip()]
    table = doc.add_table(rows=1, cols=2)
    set_table_fixed(table)
    widths = (128.0, 32.0)
    for cell, width in zip(table.rows[0].cells, widths):
        set_cell_width(cell, width)
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    left_p = table.cell(0, 0).paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if parts:
        set_run_font(left_p.add_run(parts[0]), 10.5, bold=True)
        if len(parts) > 1:
            set_run_font(left_p.add_run(" - " + " ".join(parts[1:])), 10.5)
    else:
        set_run_font(left_p.add_run(left), 10.5, bold=True)
    format_paragraph(left_p, after=0.7, line=15.3, keep_with_next=True)

    date_p = table.cell(0, 1).paragraphs[0]
    date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(date_p.add_run(date), 10.0)
    format_paragraph(date_p, after=0.7, line=15.3, keep_with_next=True)


def add_company_row(doc: Document, text: str) -> None:
    left, date = split_row(text)
    parts = [part.strip() for part in left.split("｜") if part.strip()]
    role = parts[-1] if len(parts) >= 2 else ""
    organization = "｜".join(parts[:-1]) if len(parts) >= 2 else left

    table = doc.add_table(rows=1, cols=3)
    set_table_fixed(table)
    widths = (90.0, 39.0, 31.0)
    for cell, width in zip(table.rows[0].cells, widths):
        set_cell_width(cell, width)
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    org_p = table.cell(0, 0).paragraphs[0]
    set_run_font(org_p.add_run(organization), 10.5, bold=True)
    format_paragraph(org_p, before=0.8, after=1.4, line=15.4, keep_with_next=True)

    role_p = table.cell(0, 1).paragraphs[0]
    role_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(role_p.add_run(role), 10.3)
    format_paragraph(role_p, before=0.8, after=1.4, line=15.4, keep_with_next=True)

    date_p = table.cell(0, 2).paragraphs[0]
    date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(date_p.add_run(date), 10.0)
    format_paragraph(date_p, before=0.8, after=1.4, line=15.4, keep_with_next=True)


def add_project_row(doc: Document, text: str) -> None:
    parts = [part.strip() for part in text.split("｜", 1)]
    title = parts[0]
    role = parts[1] if len(parts) > 1 else ""
    table = doc.add_table(rows=1, cols=2)
    set_table_fixed(table)
    widths = (118.0, 42.0)
    for cell, width in zip(table.rows[0].cells, widths):
        set_cell_width(cell, width)
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    title_p = table.cell(0, 0).paragraphs[0]
    set_run_font(title_p.add_run(title), 10.6, bold=True)
    format_paragraph(title_p, before=3.2, after=1.5, line=15.4, keep_with_next=True)

    role_p = table.cell(0, 1).paragraphs[0]
    role_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(role_p.add_run(role), 10.2, bold=True, color=MUTED)
    format_paragraph(role_p, before=3.2, after=1.5, line=15.4, keep_with_next=True)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(5)
    section.bottom_margin = Mm(9)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.header_distance = Mm(4)
    section.footer_distance = Mm(4)

    normal = doc.styles["Normal"]
    set_style_font(normal, 10.3)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(15.6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY

    title = doc.styles["Title"]
    set_style_font(title, 24.5, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    heading1 = doc.styles["Heading 1"]
    set_style_font(heading1, 14.4, bold=True)
    heading1.paragraph_format.space_before = Pt(0)
    heading1.paragraph_format.space_after = Pt(0)

    heading2 = doc.styles["Heading 2"]
    set_style_font(heading2, 10.6, bold=True)
    heading2.paragraph_format.space_before = Pt(0)
    heading2.paragraph_format.space_after = Pt(0)

    list_bullet = doc.styles["List Bullet"]
    set_style_font(list_bullet, 10.3)


def build(source_path: Path, output_path: Path) -> None:
    source = Document(source_path)
    source_text = [p.text for p in source.paragraphs]
    if len(source_text) < 40:
        raise ValueError(f"Unexpected source structure: {len(source_text)} paragraphs")

    doc = Document()
    configure_document(doc)
    # Remove the seed paragraph; all content is inserted deliberately.
    body = doc._body._element
    for paragraph in list(doc.paragraphs):
        body.remove(paragraph._p)

    add_header_block(doc, source_text[0], source_text[1], source_text[2])
    add_education_row(doc, source_text[3])
    add_bullet(doc, source_text[4], size=10.25, after=1.1)
    add_bullet(doc, source_text[5], size=10.25, after=1.6)
    add_education_row(doc, source_text[6])
    add_bullet(doc, source_text[7], size=10.25, after=1.1)
    add_bullet(doc, source_text[8], size=10.25, after=1.9)

    add_section_heading(doc, source_text[9])
    add_company_row(doc, source_text[10])
    add_project_row(doc, source_text[11])
    for index in range(12, 16):
        add_bullet(doc, source_text[index], size=10.25, after=1.6)
    add_project_row(doc, source_text[16])
    for index in range(17, 21):
        add_bullet(doc, source_text[index], size=10.25, after=1.6)

    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    format_paragraph(page_break, line=1.0)

    add_section_heading(doc, source_text[9], compact=True)
    add_company_row(doc, source_text[21])
    add_project_row(doc, source_text[22])
    for index in range(23, 27):
        add_bullet(doc, source_text[index], size=10.3, after=1.6)
    add_project_row(doc, source_text[27])
    for index in range(28, 32):
        add_bullet(doc, source_text[index], size=10.3, after=1.6)

    add_section_heading(doc, source_text[32])
    for index in range(33, 36):
        add_bullet(doc, source_text[index], size=10.25, after=1.5)

    add_section_heading(doc, source_text[36])
    for index in range(37, 40):
        add_bullet(doc, source_text[index], size=10.15, after=1.5)

    doc.core_properties.title = "阚海简历"
    doc.core_properties.subject = "求职简历"
    doc.core_properties.author = "阚海"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_resume_layout.py <source.docx> <output.docx>")
    build(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
