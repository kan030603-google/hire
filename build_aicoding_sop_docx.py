from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "AI-Coding通用SOP与Prompt模板.md"
DOCX_PATH = ROOT / "AI-Coding通用SOP与Prompt模板.docx"

FONT_CN = "Microsoft YaHei"
FONT_CODE = "Consolas"
BLACK = "000000"
DARK_BLUE = "2F5597"
PALE_BLUE = "EEF4FB"
LIGHT_GRAY = "D9D9D9"
CODE_BG = "F3F5F7"


def set_run_font(run, name=FONT_CN, size=None, bold=None, color=BLACK, east_asia=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia or FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        p_bdr.append(node)
    p_pr.append(p_bdr)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_CN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_CN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.22

    title = doc.styles["Title"]
    title.font.name = FONT_CN
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    title.paragraph_format.space_after = Pt(14)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 13, 14, 6),
        ("Heading 3", 11.5, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_CN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_inline_runs(paragraph, text, size=10.5):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, FONT_CODE, size - 0.5, color="333333", east_asia=FONT_CN)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, FONT_CN, size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, FONT_CN, size)


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.22
    if text.endswith((":", "：")):
        p.paragraph_format.keep_with_next = True
    add_inline_runs(p, text)
    return p


def add_list_item(doc, text, ordered=False):
    style = "List Number" if ordered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.26)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    add_inline_runs(p, text)
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = False
    p.paragraph_format.widow_control = False
    shade_paragraph(p, CODE_BG)
    for index, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, FONT_CODE, 8.5, color="202124", east_asia=FONT_CN)
        if index < len(lines) - 1:
            run.add_break()
    return p


def parse_markdown_table(lines):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    return rows


def add_table(doc, rows):
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    width_sets = {
        2: [4.8, 1.55],
        4: [0.9, 2.15, 1.65, 1.65],
        7: [1.0, 0.63, 0.7, 0.83, 0.55, 1.45, 0.8],
    }
    widths = width_sets.get(cols, [6.45 / cols] * cols)

    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if cols >= 6 and c_idx in (1, 2, 3, 4, 6):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif cols == 2 and c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            if r_idx == 0:
                set_cell_shading(cell, DARK_BLUE)
                set_run_font(run, FONT_CN, 8.5 if cols >= 6 else 9, bold=True, color="FFFFFF")
            else:
                if r_idx % 2 == 0:
                    set_cell_shading(cell, PALE_BLUE)
                set_run_font(run, FONT_CN, 8.3 if cols >= 6 else 9)
    set_repeat_table_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_cover(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    remove_paragraph_borders(p)
    p.add_run("AI Coding 通用做题 SOP 与 Agent Prompt 模板")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    r = subtitle.add_run("多候选 Rollout  MVP 筛选  Critic Actor 问题飞轮")
    set_run_font(r, FONT_CN, 12, color="444444")

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.LEFT
    summary.paragraph_format.left_indent = Inches(0.65)
    summary.paragraph_format.right_indent = Inches(0.65)
    summary.paragraph_format.line_spacing = 1.4
    summary.paragraph_format.space_after = Pt(20)
    r = summary.add_run(
        "本手册用于 README 驱动的 AI Coding 笔试。先通过四路独立 Rollout 搜索高质量初始实现，"
        "再以最高分版本为 MVP，通过 Critic、Actor 和人工评测组成可回退的问题飞轮。"
    )
    set_run_font(r, FONT_CN, 11)

    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_before = Pt(20)
    r = label.add_run("版本 1.0  2026 年 9 月")
    set_run_font(r, FONT_CN, 9.5, color="666666")
    doc.add_page_break()


def build_document():
    md = MD_PATH.read_text(encoding="utf-8")
    lines = md.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    configure_styles(doc)
    doc.core_properties.title = "AI Coding 通用做题 SOP 与 Agent Prompt 模板"
    doc.core_properties.subject = "AI Coding 笔试工作流与可复制 Agent Prompt"
    doc.core_properties.author = ""
    add_cover(doc)

    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("# "):
            i += 1
            continue

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, code_lines)
                in_code = False
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, parse_markdown_table(table_lines))
            continue

        if stripped.startswith("### "):
            p = doc.add_paragraph(stripped[4:], style="Heading 2")
            p.paragraph_format.keep_with_next = True
        elif stripped.startswith("## "):
            p = doc.add_paragraph(stripped[3:], style="Heading 1")
            p.paragraph_format.keep_with_next = True
        elif re.match(r"^-\s+", stripped):
            add_list_item(doc, re.sub(r"^-\s+", "", stripped), ordered=False)
        elif re.match(r"^\d+\.\s+", stripped):
            add_list_item(doc, re.sub(r"^\d+\.\s+", "", stripped), ordered=True)
        elif stripped:
            add_body_paragraph(doc, stripped)
        i += 1

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
