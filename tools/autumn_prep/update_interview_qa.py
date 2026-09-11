from __future__ import annotations

import io
import os
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "output" / "秋招准备材料"
QUESTION_BANK = OUTPUT / "面试八股.md"
BACKUP = ROOT / "tmp" / "autumn_prep" / "backups_before_interview_qa_20260909"


DOCS = {
    "01_项目面试_阿里_EduBot多Agent路由与任务编排.docx": {
        "section": "九、补充八股与项目映射",
        "ids": [
            "1.1", "1.2", "1.3", "1.4", "1.5", "1.7",
            "2.1", "2.2", "2.3", "2.4", "2.5", "3.4", "4.8",
            "5.1", "5.2", "5.3", "5.4", "5.5", "5.6", "5.7",
            "6.2", "6.3",
        ],
    },
    "02_项目面试_阿里_主动服务系统.docx": {
        "section": "九、补充八股与项目映射",
        "ids": ["1.1", "1.4", "1.7", "4.5", "4.6", "5.4", "6.1", "6.4", "9.3"],
    },
    "03_项目面试_高途_智能搜索对话系统.docx": {
        "section": "八、补充八股与项目映射",
        "ids": ["1.5", "1.6", "1.7", "4.1", "4.2", "4.3", "4.4", "4.5", "4.6", "4.7"],
    },
    "04_项目面试_高途_用户画像与长期记忆.docx": {
        "section": "十、补充八股与项目映射",
        "ids": ["1.4", "1.7", "4.1", "4.2", "4.3", "4.4", "4.5", "4.6", "4.7", "7.2"],
    },
    "05_项目面试_高途_多模态标签服务.docx": {
        "section": None,
        "ids": [],
    },
}


def parse_question_bank(path: Path) -> dict[str, tuple[str, str, str]]:
    text = path.read_text(encoding="utf-8")
    pattern = re.compile(
        r"^###\s+(\d+\.\d+)\s+(.+?)\n\n"
        r"\*\*回答：\*\*\s+(.+?)\n\n"
        r"\*\*项目关联：\*\*\s+(.+?)"
        r"(?=\n\n###\s+\d+\.\d+|\n\n##\s+|\Z)",
        re.MULTILINE | re.DOTALL,
    )
    questions: dict[str, tuple[str, str, str]] = {}
    for match in pattern.finditer(text):
        qid, title, answer, association = match.groups()
        questions[qid] = (
            title.strip(),
            " ".join(answer.split()),
            " ".join(association.split()),
        )
    return questions


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def update_doc(
    path: Path,
    spec: dict,
    questions: dict[str, tuple[str, str, str]],
) -> tuple[int, int]:
    # The current bytes are deliberately re-read immediately before editing so
    # every user change already present in the DOCX remains the editing baseline.
    current_bytes = path.read_bytes()
    BACKUP.mkdir(parents=True, exist_ok=True)
    backup_path = BACKUP / path.name
    if not backup_path.exists():
        backup_path.write_bytes(current_bytes)

    doc = Document(io.BytesIO(current_bytes))
    paragraphs = doc.paragraphs

    thirty = next((p for p in paragraphs if p.text.strip().startswith("30 秒版本")), None)
    two_heading = next((p for p in paragraphs if "两分钟项目陈述模板" in p.text), None)
    current_two = next((p for p in paragraphs if p.text.strip().startswith("2 分钟版本")), None)
    if thirty is None or (two_heading is None and current_two is None):
        raise RuntimeError(f"无法定位开场口播：{path.name}")

    if two_heading is not None:
        two_index = paragraphs.index(two_heading)
        if two_index + 1 >= len(paragraphs):
            raise RuntimeError(f"两分钟正文缺失：{path.name}")
        two_body = paragraphs[two_index + 1]
        two_text = two_body.text.strip()

        # Place both opening versions together near the title block.
        first_numbered_heading = next(
            p for p in paragraphs
            if p.style and p.style.name == "Heading 1" and p.text.strip().startswith("一、")
        )
        first_numbered_heading.insert_paragraph_before(f"2 分钟版本  {two_text}", style="Normal")

        # Remove the old two-minute section from the back of the document.
        remove_paragraph(two_body)
        remove_paragraph(two_heading)

    # Rebuild the generated supplement so the script is safe to rerun.
    current_paragraphs = doc.paragraphs
    source_index = next(
        i for i, p in enumerate(current_paragraphs)
        if p.style and p.style.name == "Heading 1" and p.text.strip() == "资料依据与表达边界"
    )
    supplement_index = next(
        (i for i, p in enumerate(current_paragraphs) if "补充八股与项目映射" in p.text),
        None,
    )
    if supplement_index is not None:
        for paragraph in list(current_paragraphs[supplement_index:source_index]):
            remove_paragraph(paragraph)

    source_heading = next(
        p for p in doc.paragraphs
        if p.style and p.style.name == "Heading 1" and p.text.strip() == "资料依据与表达边界"
    )

    inserted = 0
    if spec["ids"]:
        source_heading.insert_paragraph_before(spec["section"], style="Heading 1")
        source_heading.insert_paragraph_before(
            "以下问题来自通用 Agent 面试题库；回答已按本项目的事实边界筛选。通用方案不等同于当前线上能力。",
            style="Normal",
        )
        for qid in spec["ids"]:
            title, answer, association = questions[qid]
            source_heading.insert_paragraph_before(f"Q{qid}. {title}", style="Heading 2")
            source_heading.insert_paragraph_before(f"建议回答：{answer}", style="Normal")
            source_heading.insert_paragraph_before(f"项目边界：{association}", style="Normal")
            inserted += 1

    temp_path = path.with_name(path.stem + ".interview-qa.tmp.docx")
    doc.save(temp_path)
    os.replace(temp_path, path)
    return len(doc.paragraphs), inserted


def main() -> None:
    questions = parse_question_bank(QUESTION_BANK)
    expected = {qid for spec in DOCS.values() for qid in spec["ids"]}
    missing = sorted(expected - questions.keys())
    if missing:
        raise RuntimeError(f"题库缺少题号：{', '.join(missing)}")

    for filename, spec in DOCS.items():
        path = OUTPUT / filename
        paragraph_count, inserted = update_doc(path, spec, questions)
        print(f"UPDATED\t{filename}\tparagraphs={paragraph_count}\tquestions={inserted}")


if __name__ == "__main__":
    main()
